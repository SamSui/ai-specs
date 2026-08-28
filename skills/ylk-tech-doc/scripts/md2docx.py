#!/usr/bin/env python3
"""制式 md → docx 转换器。

以已有的制式 docx 为模板：保留封面、签署栏、修订记录与页眉页脚，
清空目录之后的正文，按 md 重建正文（标题 / 正文 / 表格 / 图片 / 表题图题）。

用法：
    python3 md2docx.py [--spec spec.json] \\
        <模板.docx> <源.md> <输出.docx> <封面标题行1> <封面标题行2> <文档标题>
"""
import argparse
import copy
import datetime
import os
import re
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from _docx_spec import add_spec_arg, load_spec

# 制式常量。默认值见 _docx_spec.DEFAULT_SPEC，命令行 --spec 可覆盖。
SPEC = load_spec()

PIPE_PLACEHOLDER = '\x00'   # 表格单元格内 \| 的转义占位

LIST_BULLETS = ('●', '○', '▪')   # 无序列表各层级的项目符号
LIST_INDENT_DXA = 420            # 每层列表缩进（420 dxa ≈ 2 字符）


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), v)
    return e


def strip_inline(text):
    """去除行内 markdown 标记，返回纯文本。

    链接与行内图片只保留描述文字：若原样保留，`[名称](http://…)` 会直接
    印进终稿。
    """
    text = re.sub(r'!\[([^\]]*)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'\[([^\]]*)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'`([^`]*)`', r'\1', text)
    text = re.sub(r'\*\*([^*]*)\*\*', r'\1', text)
    text = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'\1', text)
    return text.strip()


def parse_md(path):
    """把 md 解析为块序列：(kind, payload)。"""
    lines = open(path, encoding='utf-8').read().split('\n')
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        # HTML 注释块（封面信息）整体跳过
        if s.startswith('<!--'):
            while i < n and '-->' not in lines[i]:
                i += 1
            i += 1
            continue

        if not s:
            i += 1
            continue

        # 围栏代码块：逐行原样保留，不作行内标记处理。
        # docx 侧用等宽段落呈现——若用编号列表，Word 会吞掉行首的
        # `ur.`、`id` 等前缀文本。
        m = re.match(r'^(`{3,}|~{3,})(.*)$', s)
        if m:
            fence = m.group(1)[0] * 3
            lang = m.group(2).strip().lower()
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith(fence):
                code.append(lines[i].rstrip())
                i += 1
            i += 1                      # 跳过收尾围栏
            # mermaid 是绘图源码而非文档内容：docx 应由插图版 md 生成，
            # 此处整块丢弃并告警，避免把图源码当正文贴进终稿。
            if lang == 'mermaid':
                sys.stderr.write(
                    '警告：跳过 mermaid 代码块（第 %d 行）。'
                    'docx 应以插图版 md 为输入。\n' % (i - len(code) - 1))
            elif code:
                blocks.append(('code', code))
            continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            blocks.append(('h%d' % len(m.group(1)), strip_inline(m.group(2))))
            i += 1
            continue

        # 图片
        m = re.match(r'^!\[[^\]]*\]\(([^)]+)\)$', s)
        if m:
            blocks.append(('img', m.group(1)))
            i += 1
            continue

        # 表格
        if s.startswith('|') and i + 1 < n and re.match(r'^\|[\s:|-]+\|$', lines[i + 1].strip()):
            rows = []
            sep_idx = i + 1     # 分隔行只有紧跟表头的这一行
            while i < n and lines[i].strip().startswith('|'):
                if i != sep_idx:
                    # 数据行也可能长得像分隔行（如整行都是「-」表示不适用），
                    # 故按下标判定而非按内容，否则该行会被静默吞掉。
                    raw = lines[i].strip().replace('\\|', PIPE_PLACEHOLDER)
                    cells = [strip_inline(c).replace(PIPE_PLACEHOLDER, '|')
                             for c in raw.strip('|').split('|')]
                    rows.append(cells)
                i += 1
            blocks.append(('table', rows))
            continue

        # 表题 / 图题（**表 X-Y ...** 或 **图 X-Y ...**）
        m = re.match(r'^\*\*((?:表|图)\s*\d+-\d+.*?)\*\*$', s)
        if m:
            blocks.append(('caption', m.group(1)))
            i += 1
            continue

        # 引用块
        if s.startswith('>'):
            blocks.append(('quote', strip_inline(s.lstrip('> '))))
            i += 1
            continue

        # 列表：整块收集后用缩进栈换算层级。
        # 缩进步长（2 空格 / 4 空格）因文档而异，不能按固定除数换算；
        # 也不能把全块缩进量排序取下标——相邻两个列表的步长若不同
        # （如无序用 2 空格、有序用 3 空格），排序会把层级算串。
        if re.match(r'^([-*]|\d+[.)])\s+', s):
            raw_items = []
            while i < n:
                cur = lines[i]
                t = cur.strip()
                if not t:                   # 列表项之间允许空行
                    if i + 1 < n and re.match(r'^([-*]|\d+[.)])\s+',
                                              lines[i + 1].strip()):
                        i += 1
                        continue
                    break
                m = re.match(r'^([-*]|\d+[.)])\s+(.*)$', t)
                if not m:
                    break
                indent = len(cur) - len(cur.lstrip(' \t'))
                marker = m.group(1) if m.group(1) not in ('-', '*') else ''
                raw_items.append((indent, marker, strip_inline(m.group(2))))
                i += 1
            stack = []          # 各层级的缩进量，单调递增
            for indent, marker, text in raw_items:
                while stack and indent < stack[-1]:
                    stack.pop()
                if not stack or indent > stack[-1]:
                    stack.append(indent)
                blocks.append(('li', (len(stack) - 1, marker, text)))
            continue

        blocks.append(('p', strip_inline(s)))
        i += 1
    return blocks


def clear_body_after(doc, keep_upto_index):
    """删除 body 中下标 > keep_upto_index 的元素，保留末尾 sectPr。"""
    body = doc.element.body
    children = list(body.iterchildren())
    sect = children[-1] if children[-1].tag == qn('w:sectPr') else None
    for ch in children[keep_upto_index + 1:]:
        if ch is sect:
            continue
        body.remove(ch)
    return sect


def _is_toc_style(name):
    n = name.lower()
    return n.startswith('toc') or n.startswith('目录') or n.startswith('目次')


def find_toc_end(doc):
    """返回目录最后一个元素的下标。

    取**第一段连续**的目录段落的末尾，而非全局最后一个 toc 段落：模板正文区
    若另有图表目录，全局取最后会让 clear_body_after 几乎不清空旧正文。
    """
    from docx.text.paragraph import Paragraph
    body = doc.element.body
    last = -1
    for idx, ch in enumerate(body.iterchildren()):
        if ch.tag != qn('w:p'):
            continue
        p = Paragraph(ch, doc)
        if _is_toc_style(p.style.name):
            last = idx
        elif last >= 0 and p.text.strip():
            break       # 已越过目录区（空段落不算中断）
    return last


# CT_TblPrBase 的 schema 序列。tblPr 子元素必须按此顺序，否则 Word 拒绝打开。
TBLPR_ORDER = ('w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual',
               'w:tblStyleRowBandSize', 'w:tblStyleColBandSize', 'w:tblW',
               'w:jc', 'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd',
               'w:tblLayout', 'w:tblCellMar', 'w:tblLook')


def _tblpr_put(tblPr, el):
    """按 schema 顺序插入 tblPr 子元素，已存在的同名元素先移除。"""
    tag = el.tag
    old = tblPr.find(tag)
    if old is not None:
        tblPr.remove(old)
    rank = TBLPR_ORDER.index('w:' + tag.split('}')[-1])
    for ch in tblPr:
        name = 'w:' + ch.tag.split('}')[-1]
        if name in TBLPR_ORDER and TBLPR_ORDER.index(name) > rank:
            ch.addprevious(el)
            return
    tblPr.append(el)


def format_table(doc, t):
    """按制式统一表格版式：等宽居中、全边框、表头底纹与加粗、单元格紧凑段距。"""
    tbl = t._tbl
    tblPr = tbl.find(qn('w:tblPr'))

    old = tblPr.find(qn('w:tblLook'))
    if old is not None:
        tblPr.remove(old)

    _tblpr_put(tblPr, _el('w:tblW', w=str(SPEC['tbl_width_dxa']), type='dxa'))
    _tblpr_put(tblPr, _el('w:jc', val='center'))

    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(_el('w:' + side, val='single', color='auto', sz=SPEC['border_sz'], space='0'))
    _tblpr_put(tblPr, borders)

    _tblpr_put(tblPr, _el('w:tblLayout', type='autofit'))
    cellmar = OxmlElement('w:tblCellMar')
    for side, w in (('top', '0'), ('left', '108'), ('bottom', '0'), ('right', '108')):
        cellmar.append(_el('w:' + side, w=w, type='dxa'))
    _tblpr_put(tblPr, cellmar)

    # 列宽均分至总宽，避免 autofit 下列宽失衡
    ncol = len(t.columns)
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement('w:tblGrid')
    per = SPEC['tbl_width_dxa'] // ncol
    for ci in range(ncol):
        w = SPEC['tbl_width_dxa'] - per * (ncol - 1) if ci == ncol - 1 else per
        grid.append(_el('w:gridCol', w=str(w)))
    tbl.insert(list(tbl).index(tblPr) + 1, grid)
    compact = _pick_style(doc, 'Compact')

    for ri, row in enumerate(t.rows):
        tr = row._tr
        trPr = tr.find(qn('w:trPr'))
        if trPr is None:
            trPr = OxmlElement('w:trPr')
            tr.insert(0, trPr)
        for tag in ('w:trHeight', 'w:tblHeader', 'w:jc'):
            old = trPr.find(qn(tag))
            if old is not None:
                trPr.remove(old)
        trPr.append(_el('w:trHeight', val='0', hRule='atLeast'))
        if ri == 0:
            trPr.append(OxmlElement('w:tblHeader'))   # 表头跨页重复
        trPr.append(_el('w:jc', val='center'))

        for cell in row.cells:
            tcPr = cell._tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = OxmlElement('w:tcPr')
                cell._tc.insert(0, tcPr)
            for tag in ('w:shd', 'w:vAlign'):
                old = tcPr.find(qn(tag))
                if old is not None:
                    tcPr.remove(old)
            if ri == 0:
                tcPr.append(_el('w:shd', val='clear', color='auto', fill=SPEC['header_fill']))
            tcPr.append(_el('w:vAlign', val='center'))

            for para in cell.paragraphs:
                para.style = compact
                pPr = para._p.get_or_add_pPr()
                for tag in ('w:snapToGrid', 'w:spacing', 'w:ind', 'w:jc'):
                    old = pPr.find(qn(tag))
                    if old is not None:
                        pPr.remove(old)
                pPr.append(_el('w:snapToGrid', val='0'))
                pPr.append(_el('w:spacing', line='240', lineRule='auto'))
                pPr.append(_el('w:ind', left='0', leftChars='0', right='0',
                               rightChars='0', firstLine='0', firstLineChars='0'))
                pPr.append(_el('w:jc', val='center' if ri == 0 else 'left'))
                for run in para.runs:
                    run.bold = True if ri == 0 else None


def _toc_rpr():
    """目录条目字体，取自制式（默认宋体小四 = sz 24）。"""
    rPr = OxmlElement('w:rPr')
    font = SPEC['toc_font']
    rPr.append(_el('w:rFonts', ascii=font, hAnsi=font,
                   eastAsia=font, cs=SPEC['toc_font_cs']))
    rPr.append(_el('w:kern', val=SPEC['toc_kern']))
    rPr.append(_el('w:sz', val=SPEC['toc_sz']))
    rPr.append(_el('w:szCs', val=SPEC['toc_sz']))
    return rPr


def _field_run(kind, dirty=False):
    """域字符 run。dirty=True 时标记该域为待重算，打开文档即刷新页码。"""
    r = OxmlElement('w:r')
    r.append(_toc_rpr())
    fc = _el('w:fldChar', fldCharType=kind)
    if dirty:
        fc.set(qn('w:dirty'), 'true')
    r.append(fc)
    return r


def _instr_run(text):
    r = OxmlElement('w:r')
    r.append(_toc_rpr())
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = text
    r.append(it)
    return r


def _text_run(text):
    r = OxmlElement('w:r')
    r.append(_toc_rpr())
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return r


def add_bookmark(p, name, bid):
    """给标题段落加 _Toc 书签，供目录超链接与页码域引用。"""
    start = _el('w:bookmarkStart', id=str(bid), name=name)
    end = _el('w:bookmarkEnd', id=str(bid))
    p._p.insert(0, start)
    p._p.append(end)


def add_page_break_before(p):
    """在段落前插入分页符（章首起新页）。"""
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:pageBreakBefore')) is None:
        pPr.append(OxmlElement('w:pageBreakBefore'))


def add_section_break(doc, template_sect):
    """在正文起始处插入分节符（复用模板的页脚与页码设置）。"""
    # 用 Normal 而非 Title：Title 样式常带 outlineLvl，Word 重建目录时
    # 会把这个空段落收成一条空目录项。
    p = doc.add_paragraph(style=_pick_style(doc, 'Normal'))
    pPr = p._p.get_or_add_pPr()
    pPr.append(copy.deepcopy(template_sect))
    return p


def rebuild_toc(doc, headings):
    """按新正文的标题重建目录条目文本，页码交由 Word 更新域时回填。"""
    from docx.text.paragraph import Paragraph
    body = doc.element.body
    toc_ps = [ch for ch in body.iterchildren()
              if ch.tag == qn('w:p')
              and Paragraph(ch, doc).style.name.lower().startswith('toc')]
    if not toc_ps:
        # find_toc_end 认中文「目录/目次」样式，此处只认英文 toc 样式（条目段落
        # 的样式必须是 toc N，否则重写出来的条目 Word 不认）。模板若只有中文
        # 标题样式而无 toc 条目样式，目录就无从重建——静默返回会让终稿留着
        # 旧文档的目录，故告警。
        sys.stderr.write('警告：模板中未找到 toc 样式的目录条目段落，'
                         '目录未重建——请确认模板目录由 TOC 域生成。\n')
        return
    # 首个 toc 段落承载 TOC 域代码，保留域、剥离旧域结果文本；其余按新标题重写
    keeper = toc_ps[0]
    for ch in toc_ps[1:]:
        body.remove(ch)
    # 只保留最外层 TOC 域的 begin / instrText / separate 三个 run，
    # 其后的嵌套 HYPERLINK 域与旧域结果文本一律删除
    kept = 0
    for r in list(keeper.findall(qn('w:r'))):
        instr = r.find(qn('w:instrText'))
        fld = r.find(qn('w:fldChar'))
        typ = fld.get(qn('w:fldCharType')) if fld is not None else None
        if kept == 0 and typ == 'begin':
            fld.set(qn('w:dirty'), 'true')   # 打开文档即重建整个目录
            kept = 1
        elif kept == 1 and instr is not None and 'TOC' in (instr.text or ''):
            kept = 2
        elif kept == 2 and typ == 'separate':
            kept = 3
        else:
            keeper.remove(r)
    for hl in keeper.findall(qn('w:hyperlink')):
        keeper.remove(hl)
    anchor = keeper
    for lvl, text, bm in headings:
        p = copy.deepcopy(keeper)
        for r in p.findall(qn('w:r')):
            p.remove(r)
        pPr = p.find(qn('w:pPr'))
        if pPr is not None:
            st = pPr.find(qn('w:pStyle'))
            if st is not None:
                # 模板可能只定义到 toc 1；缺失时沿用上一级，避免 KeyError
                for cand in range(lvl, 0, -1):
                    try:
                        st.set(qn('w:val'), doc.styles['toc %d' % cand].style_id)
                        break
                    except KeyError:
                        continue
        # 条目结构与参考一致：HYPERLINK 域包裹「标题 + 点线制表 + PAGEREF 页码域」
        p.append(_field_run('begin'))
        p.append(_instr_run(' HYPERLINK \\l %s ' % bm))
        p.append(_field_run('separate'))
        p.append(_text_run(text))
        tab_run = OxmlElement('w:r')
        tab_run.append(OxmlElement('w:tab'))
        p.append(tab_run)
        p.append(_field_run('begin', dirty=True))
        p.append(_instr_run(' PAGEREF %s \\h ' % bm))
        p.append(_field_run('separate'))
        p.append(_text_run('1'))     # 占位页码，打开文档时由 dirty 域重算
        p.append(_field_run('end'))
        p.append(_field_run('end'))  # 闭合 HYPERLINK 域
        anchor.addnext(p)
        anchor = p

    # TOC 域需 begin/end 配对：end 置于最后一个条目末尾
    anchor.append(_field_run('end'))

    # 清理目录区内的游离域字符：部分模板把 TOC 域的收尾 fldChar 放在
    # 目录之后的非 toc 样式段落里（如空段），上面按 toc 样式收集时收不到。
    # 若不删，与新补的 end 叠加会导致 begin/end 失配，Word 报「未找到引用源」。
    tail = anchor.getnext()
    while tail is not None:
        nxt = tail.getnext()
        if tail.tag == qn('w:p'):
            style = Paragraph(tail, doc).style.name
            # 判据须与上方 toc_ps 的收集判据一致：只有被 rebuild 重写过的
            # 条目段落才放行。「目录 N」这类标题样式不在其列，其中的游离域
            # 字符正是要清掉的对象。
            if style.lower().startswith('toc'):
                tail = nxt
                continue
            # 正文起点即停：正文首段若自带域（页码引用、SEQ 等），不能误删
            if style.startswith('Heading') or style == 'Title':
                break
            stray = [r for r in tail.findall(qn('w:r'))
                     if r.find(qn('w:fldChar')) is not None
                     or r.find(qn('w:instrText')) is not None]
            if not stray:
                break          # 已进入正文区，停止清理
            for r in stray:
                tail.remove(r)
        tail = nxt


def mark_update_fields(doc):
    """打开文档时提示更新域，使目录页码自动重算。"""
    settings = doc.settings.element
    if settings.find(qn('w:updateFields')) is None:
        settings.append(_el('w:updateFields', val='true'))


def _pick_style(doc, name):
    """返回可用的样式名。模板未定义该级标题时逐级降级，避免直接崩溃。

    模板的标题层级深度不一（有的只到 Heading 2），md 中的更深层级
    退化为上一级样式呈现，正文内容不丢。
    """
    fallback = {'Heading 3': ['Heading 2', 'Heading 1'],
                'Heading 2': ['Heading 1'],
                'Body Text': ['Normal'],
                'Caption': ['Body Text', 'Normal'],
                'Compact': ['Body Text', 'Normal'],
                'Title': ['Heading 1']}
    for cand in [name] + fallback.get(name, []):
        try:
            doc.styles[cand]
            return cand
        except KeyError:
            continue
    return 'Normal'


def _set_para_text(p, text):
    """把段落改写为单个 run 的指定文本，保留首个 run 的字体格式。"""
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text


def set_cover(doc, line1, line2, cover_date=None):
    """改写封面：项目名行、文档类型行、日期栏。

    前提（不满足则告警并跳过对应改写）：
    - 封面标题行是 body 顶部的**普通段落**，不在表格里；
    - 标题行文本不含冒号（冒号用于区分其后的「项目令号：」等信息栏）；
    - 日期栏以 SPEC['date_field'] 开头。

    取信息栏之前的最后两个无冒号段落作为项目名行与文档类型行，
    因此不依赖具体项目名，模板若无单位名行也成立。
    """
    from docx.text.paragraph import Paragraph
    body = doc.element.body
    ps = [Paragraph(ch, doc) for ch in list(body.iterchildren())[:SPEC['cover_scan']]
          if ch.tag == qn('w:p')]
    if cover_date is None:
        cover_date = datetime.date.today().strftime(SPEC['date_format'])

    title_lines = []
    seen_info = False
    date_set = False
    for p in ps:
        t = p.text.strip()
        if not t:
            continue
        if '：' in t or ':' in t:      # 信息栏，标题行到此为止
            seen_info = True
            if t.startswith(SPEC['date_field']):
                _set_para_text(p, '%s：     %s' % (SPEC['date_field'], cover_date))
                date_set = True
            continue
        if not seen_info:
            title_lines.append(p)

    if len(title_lines) < 2:
        sys.stderr.write('警告：封面未找到两行标题段落（找到 %d 行），'
                         '封面标题未改写——模板封面可能在表格中。\n'
                         % len(title_lines))
    if not date_set:
        sys.stderr.write('警告：封面未找到「%s」栏，日期未改写。\n' % SPEC['date_field'])

    for p, text in zip(title_lines[-2:], (line1, line2)):
        _set_para_text(p, text)


def build(template, md_path, out_path, cover1, cover2, doc_title):
    doc = Document(template)
    toc_end = find_toc_end(doc)
    if toc_end <= 0:
        sys.exit('模板中未找到目录：请确认模板含 toc 样式段落')

    # 取模板正文节的 sectPr（含页脚与页码重编），用于目录后的分节
    body_sect = None
    from docx.text.paragraph import Paragraph
    for ch in list(doc.element.body.iterchildren())[toc_end + 1:]:
        if ch.tag != qn('w:p'):
            continue
        pPr = ch.find(qn('w:pPr'))
        if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
            body_sect = copy.deepcopy(pPr.find(qn('w:sectPr')))
            break

    sect = clear_body_after(doc, toc_end)
    set_cover(doc, cover1, cover2)

    blocks = parse_md(md_path)
    base = os.path.dirname(os.path.abspath(md_path))

    # 层级约定：`#` 是文档标题，`##` 起才是章。多个 `#` 说明源文档用的是
    # 另一套层级，此时静默丢弃会让整篇章标题消失，故直接报错。
    if sum(1 for k, _ in blocks if k == 'h1') > 1:
        sys.exit('源 md 含多个 `#` 一级标题。本转换器约定 `#` 为文档标题、'
                 '`##` 为章标题，请调整层级后重试。')

    # 目录与正文之间分节，正文页码自第 1 页重编
    if body_sect is not None:
        add_section_break(doc, body_sect)

    headings = []           # 供目录重建
    h1_count = 0            # 一级标题序号（Word 自动编号，此处仅用于目录文本）
    h2_count = 0
    title_added = False
    bm_id = 1000            # 书签编号，避开模板既有书签
    for kind, payload in blocks:
        if kind == 'h1':
            if not title_added:
                p = doc.add_paragraph(doc_title, style=_pick_style(doc, 'Title'))
                bm_id += 1
                bm = '_Toc%d' % bm_id
                add_bookmark(p, bm, bm_id)
                headings.append((1, doc_title, bm))
                title_added = True
            continue
        if kind in ('h2', 'h3', 'h4', 'h5', 'h6'):
            # h5/h6 无对应样式，压到 Heading 3 呈现，不进目录
            lvl = {'h2': 'Heading 1', 'h3': 'Heading 2', 'h4': 'Heading 3',
                   'h5': 'Heading 3', 'h6': 'Heading 3'}[kind]
            # 去掉 md 中的编号前缀，交给 Word 自动编号
            text = re.sub(r'^[\d.]+\s*', '', payload)
            p = doc.add_paragraph(text, style=_pick_style(doc, lvl))
            if kind in ('h2', 'h3'):
                bm_id += 1
                bm = '_Toc%d' % bm_id
                add_bookmark(p, bm, bm_id)
            if kind == 'h2':
                h1_count += 1
                h2_count = 0
                # 每章一级标题从新页开始（首章紧随文档标题，不另起页）
                if h1_count > 1:
                    add_page_break_before(p)
                headings.append((1, '%d. %s' % (h1_count, text), bm))
            elif kind == 'h3':
                h2_count += 1
                headings.append((2, '%d.%d %s' % (h1_count, h2_count, text), bm))
        elif kind == 'p':
            doc.add_paragraph(payload, style=_pick_style(doc, 'Body Text'))
        elif kind == 'quote':
            doc.add_paragraph(payload, style=_pick_style(doc, 'Body Text'))
        elif kind == 'li':
            depth, marker, text = payload
            # 项目符号手工写进文本：Word 的自动编号列表会吞掉行首前缀，
            # 且样式随模板漂移，制式文档里不如固定符号可控。
            prefix = marker or LIST_BULLETS[min(depth, len(LIST_BULLETS) - 1)]
            p = doc.add_paragraph('%s %s' % (prefix, text),
                                  style=_pick_style(doc, 'Body Text'))
            ind = str(LIST_INDENT_DXA * (depth + 1))
            pPr = p._p.get_or_add_pPr()
            pPr.append(_el('w:ind', left=ind, hanging=str(LIST_INDENT_DXA),
                           firstLine='0', firstLineChars='0'))
        elif kind == 'code':
            # 等宽段落逐行呈现，保留缩进；不用编号列表（Word 会吞行首前缀）
            for ln in payload:
                p = doc.add_paragraph(style=_pick_style(doc, 'Body Text'))
                pPr = p._p.get_or_add_pPr()
                pPr.append(_el('w:ind', left='0', leftChars='0',
                               firstLine='0', firstLineChars='0'))
                r = p.add_run(ln.replace('\t', '    ') or ' ')
                r.font.name = SPEC['code_font']
                r._element.rPr.rFonts.set(qn('w:eastAsia'), SPEC['code_font'])
                r._element.rPr.rFonts.set(qn('w:cs'), SPEC.get('code_font_cs', SPEC['code_font']))
                sz = str(int(SPEC.get('body_sz', '24')) + int(SPEC.get('code_sz_delta', -2)))
                r._element.rPr.append(_el('w:sz', val=sz))
                r._element.rPr.append(_el('w:szCs', val=sz))
                r._element.rPr.append(_el('w:color', val=SPEC.get('code_color', '595959')))
        elif kind == 'caption':
            p = doc.add_paragraph(payload, style=_pick_style(doc, 'Caption'))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == 'img':
            # 图片段落沿用 Body Text 样式并居中、去首行缩进
            p = doc.add_paragraph(style=_pick_style(doc, 'Body Text'))
            pPr = p._p.get_or_add_pPr()
            pPr.append(_el('w:ind', left='0', leftChars='0',
                           firstLine='0', firstLineChars='0'))
            pPr.append(_el('w:jc', val='center'))
            p.add_run().add_picture(os.path.join(base, payload), width=Inches(SPEC['img_max_width_in']))
        elif kind == 'table':
            rows = payload
            ncol = max((len(r) for r in rows), default=0)
            if not rows or ncol == 0:
                sys.stderr.write('警告：跳过空表格。\n')
                continue
            t = doc.add_table(rows=len(rows), cols=ncol)
            try:
                t.style = SPEC['table_style']
            except KeyError:
                t.style = 'Table Grid'
                sys.stderr.write('警告：模板无 %s 样式，回退 Table Grid，'
                                 % SPEC['table_style'] +
                                 '表格字体与段距可能与模板不一致。\n')
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    t.cell(ri, ci).text = row[ci] if ci < len(row) else ''
            format_table(doc, t)

    rebuild_toc(doc, headings)
    mark_update_fields(doc)
    if sect is not None:
        doc.element.body.append(sect)
    doc.save(out_path)
    print('已生成', out_path)


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('template', help='样板 docx，提供封面/签署页/目录/样式集')
    ap.add_argument('md', help='源 markdown（须为插图版，非 mermaid 版）')
    ap.add_argument('out', help='输出 docx')
    ap.add_argument('cover1', help='封面标题行 1（项目名）')
    ap.add_argument('cover2', help='封面标题行 2（文档类型）')
    ap.add_argument('title', help='正文首页的文档标题')
    add_spec_arg(ap)
    a = ap.parse_args()

    SPEC = load_spec(a.spec)
    for path in (a.template, a.md):
        if not os.path.isfile(path):
            sys.exit('文件不存在：%s' % path)
    if os.path.abspath(a.out) == os.path.abspath(a.template):
        sys.exit('输出路径与模板相同，会覆盖模板：%s' % a.out)
    build(a.template, a.md, a.out, a.cover1, a.cover2, a.title)
